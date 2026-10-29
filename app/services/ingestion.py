import hashlib
from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument
from app.core.config import get_settings
from app.rag.common.vectorstore import add_documents, delete_vectors
from app.services.documents import PUBLIC_ACCESS, get_document, register_document, unregister_document

settings = get_settings()

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))


def chunk_ids(domain: str, filename: str, count: int, start: int = 0) -> list[str]:
    """Stable vector ids: re-indexing a file overwrites its vectors instead of duplicating them."""
    key = hashlib.sha1(f"{domain}/{filename}".encode("utf-8")).hexdigest()[:16]
    return [f"{domain}-{key}-{i}" for i in range(start, count)]


def ingest_file(path: Path, domain: str, namespace: str, access: list[str] | None = None) -> int:
    """Load, tag, chunk and index one file into the domain's namespace. Returns the chunk count.
    `access` lists the groups allowed to retrieve it; every chunk carries it for the filter."""
    access = sorted(set(access or PUBLIC_ACCESS))
    docs = load_file(path)
    for doc in docs:
        doc.metadata.update(domain=domain, filename=path.name, access=access)
    chunks = [c for c in chunk_documents(docs) if c.page_content.strip()]
    if not chunks:
        raise ValueError(f"No text could be extracted from {path.name}")

    add_documents(chunks, namespace, ids=chunk_ids(domain, path.name, len(chunks)))
    previous = get_document(domain, path.name)
    if previous and previous["chunks"] > len(chunks):
        # The new version is shorter: drop the leftover vectors of the old one.
        delete_vectors(chunk_ids(domain, path.name, previous["chunks"], start=len(chunks)), namespace)
    register_document(domain, path.name, len(chunks), access)
    return len(chunks)


def ingest_files(paths: Iterable[Path], domain: str, namespace: str,
                 access: dict[str, list[str]] | None = None) -> int:
    """`access` maps a file name to its groups; files not listed are public."""
    return sum(ingest_file(path, domain, namespace, (access or {}).get(path.name)) for path in paths)


def delete_document(domain: str, filename: str, namespace: str) -> None:
    """Remove a document's vectors, registry entry and uploaded copy. Raise KeyError if unknown."""
    doc = get_document(domain, filename)
    if doc is None:
        raise KeyError(filename)
    delete_vectors(chunk_ids(domain, filename, doc["chunks"]), namespace)
    unregister_document(domain, filename)
    (Path(settings.upload_dir) / domain / filename).unlink(missing_ok=True)
